"""Render report_content.py into the VIT B.Tech project report template.

Keeps the template's front matter, styles, chapter numbering ("Chapter N"), caption styles and TOC fields; replaces the
placeholder text and regenerates every chapter. Output: TruthLens_Final_Project_Report.docx (fields are refreshed by
refresh_fields.ps1, or by Word on opening).

    python make_figures.py && python build_report.py
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT / "docs" / "paper"))
import report_content as RC  # noqa: E402
from paper_content import REFERENCES as PAPER_REFS  # noqa: E402

TEMPLATE = HERE / "work" / "template.docx"
OUT = HERE / "TruthLens_Final_Project_Report.docx"
CROPS = HERE / "work" / "crops"
CROPS.mkdir(parents=True, exist_ok=True)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


# ------------------------------------------------------------------------------------------- references
def latex_to_text(s: str) -> str:
    for a, b in [(r'\"o', "ö"), (r'\"u', "ü"), (r'\"a', "ä"), (r"{\ss}", "ß"), ("``", "“"), ("''", "”"), ("--", "–"),
                 (r"\&", "&"), ("~", " ")]:
        s = s.replace(a, b)
    s = re.sub(r"\\emph\{([^}]*)\}", r"\1", s)
    return s.replace("{", "").replace("}", "")


REFS = {k: latex_to_text(v) for k, v in PAPER_REFS.items()}
REFS.update(RC.EXTRA_REFERENCES)
ORDER: list[str] = []
CITE = re.compile(r"\{([a-z0-9_,]+)\}")


def collect(text: str):
    for m in CITE.finditer(text):
        keys = m.group(1).split(",")
        if not all(k in REFS for k in keys):
            continue  # e.g. an API path such as /detect/{id}
        for k in keys:
            if k not in ORDER:
                ORDER.append(k)


def cite(text: str) -> str:
    def rep(m):
        if not all(k in REFS for k in m.group(1).split(",")):
            return m.group(0)
        return "[" + ", ".join(str(ORDER.index(k) + 1) for k in m.group(1).split(",")) + "]"
    return CITE.sub(rep, text)


def walk_text(block):
    kind = block[0]
    if kind in ("p", "h1", "h2", "h3", "ah"):
        yield block[1]
    elif kind in ("bul", "num"):
        yield from block[1]
    elif kind == "fig":
        yield block[2]
    elif kind == "tab":
        yield block[1]
        yield from block[2]
        for r in block[3]:
            yield from r


ALL_BLOCKS = RC.BODY + getattr(RC, "APPENDICES", [])
for b in ALL_BLOCKS:
    for t in walk_text(b):
        collect(t)

# figure numbers ("chapter.n") by key, so the text can say {fig:key} instead of a hard-coded number
FIG_NO: dict[str, str] = {}
_ch, _lab, _fn = 0, "", 0
for b in ALL_BLOCKS:
    if b[0] == "h1":
        _ch, _fn = _ch + 1, 0
        _lab = str(_ch)
    elif b[0] == "app":
        _lab, _fn = b[1], 0
    elif b[0] == "fig":
        _fn += 1
        FIG_NO[b[1]] = f"{_lab}.{_fn}"
FIGREF = re.compile(r"\{fig:([@a-z0-9_]+)\}")


def figref(text: str) -> str:
    return FIGREF.sub(lambda m: FIG_NO[m.group(1)], text)


# ------------------------------------------------------------------------------------------- run helpers
def add_runs(p, text: str, size: int | None = None, bold_all=False):
    """**bold** and *italic* markup -> runs."""
    text = cite(figref(text))
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            r = p.add_run(tok)
        if bold_all:
            r.bold = True
        if size:
            r.font.size = docx.shared.Pt(size)


def set_ppr(p, xml: str):
    old = p._p.find(qn("w:pPr"))
    if old is not None:
        p._p.remove(old)
    p._p.insert(0, parse_xml(f"<w:pPr {nsdecls('w')}>{xml}</w:pPr>"))


def set_text_keep_format(p, text: str):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


# ------------------------------------------------------------------------------------------- document setup
doc = docx.Document(str(TEMPLATE))
body = doc.element.body
P = doc.paragraphs


def find(text: str, start: int = 0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if text in p.text:
            return i, p
    raise KeyError(text)


tpl_tblpr = copy.deepcopy(doc.tables[1]._tbl.find(qn("w:tblPr")))
tpl_hdr_tcpr = copy.deepcopy(doc.tables[1].rows[0].cells[0]._tc.find(qn("w:tcPr")))
tpl_body_tcpr = copy.deepcopy(doc.tables[1].rows[1].cells[0]._tc.find(qn("w:tcPr")))
_, ref_tpl = find("D. Xue, Y. Q. Chen")
ref_ppr = copy.deepcopy(ref_tpl._p.find(qn("w:pPr")))

# --- front matter
names = [f"{n} ({r})" for n, r in RC.STUDENTS]
plain_names = [n for n, _ in RC.STUDENTS]
set_text_keep_format(find("Title of the Project")[1], RC.TITLE)
i0, _ = find("Surname Firstname")
for k in range(4):
    set_text_keep_format(doc.paragraphs[i0 + k], names[k])
set_text_keep_format(find("Dr. / Prof. Supervisor Name")[1], RC.GUIDE)
set_text_keep_format(find("2025-2026")[1], RC.YEAR)
# the three-line title pushes the cover onto a second page: drop two of the empty spacer paragraphs above the address
iD, _ = find("Department of Information Technology")
spacers = [p for p in doc.paragraphs[iD - 5:iD] if not p.text.strip() and not p._p.findall(".//" + qn("w:drawing"))]
for p in spacers[:2]:
    p._p.getparent().remove(p._p)
set_text_keep_format(find("“Title of the Project”")[1], f"“{RC.TITLE}”")
i1, _ = find("Surname Firstname", i0 + 4)
for k in range(4):
    set_text_keep_format(doc.paragraphs[i1 + k], names[k])
_, appr = find("This is to certify that the dissertation entitled")
for r in appr.runs:
    r.text = (r.text.replace("Title of the Project", RC.TITLE)
              .replace("Student Names", ", ".join(plain_names[:-1]) + " and " + plain_names[-1])
              .replace("Dr. / Prof. Supervisor Name", RC.GUIDE))

_, ack = find("[Express your gratitude")
anchor = ack
set_text_keep_format(ack, "")
add_runs(ack, RC.ACK[0])
ack.runs[0]._r.getparent().remove(ack.runs[0]._r)  # the emptied original run
for t in RC.ACK[1:]:
    new = copy.deepcopy(ack._p)
    for r in new.findall(qn("w:r")):
        new.remove(r)
    anchor._p.addnext(new)
    anchor = docx.text.paragraph.Paragraph(new, ack._parent)
    add_runs(anchor, t)
iS, _ = find("Student Name 1 (Roll No.)")
for k in range(4):
    set_text_keep_format(doc.paragraphs[iS + k], names[k])

set_text_keep_format(find("[Write the abstract here")[1], RC.ABSTRACT)
_, kw = find("Keywords:")
kw.runs[-1].text = RC.KEYWORDS
for r in kw.runs[1:-1]:
    r.text = ""

_, instr = find("[List abbreviations")
iA, first_abbr = find("Artificial Intelligence")
abbr_tpl = copy.deepcopy(first_abbr._p)
old_abbr = [p for p in doc.paragraphs[iA:iA + 12] if "–" in p.text]
anchor_el = old_abbr[-1]._p
for ab, full in reversed(RC.ABBREVIATIONS):
    e = copy.deepcopy(abbr_tpl)
    rs = e.findall(qn("w:r"))
    rs[0].find(qn("w:t")).text = ab
    for r in rs[2:]:
        e.remove(r)
    e.append(parse_xml(f'<w:r {nsdecls("w")}><w:t xml:space="preserve">–  {full}</w:t></w:r>'))
    for pe in e.findall(qn("w:proofErr")):
        e.remove(pe)
    anchor_el.addnext(e)
for p in old_abbr + [instr]:
    p._p.getparent().remove(p._p)

# --- remove the template chapters (from the first Heading 1 to the end, keeping the final sectPr)
first_h1 = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
el = first_h1._p
while el is not None and el.tag != qn("w:sectPr"):
    nxt = el.getnext()
    body.remove(el)
    el = nxt

# ------------------------------------------------------------------------------------------- block renderers
numbering = doc.part.numbering_part.element
chapter = 0
fig_n = tab_n = 0


def new_num(abstract_id: int) -> int:
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    nid = max(ids) + 1
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{abstract_id}"/>'
        f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
    return nid


label = ""  # caption prefix: chapter number, or the appendix letter


def heading(level: int, text: str):
    global chapter, fig_n, tab_n, label
    p = doc.add_paragraph(style=f"Heading {level}")
    set_ppr(p, f'<w:pStyle w:val="Heading{level}"/><w:numPr><w:ilvl w:val="{level - 1}"/><w:numId w:val="2"/></w:numPr>')
    if level == 1:
        chapter += 1
        label = str(chapter)
        fig_n = tab_n = 0
        p.add_run().add_break()
    p.add_run(text)


def para(text: str):
    p = doc.add_paragraph()
    set_ppr(p, '<w:spacing w:after="120"/><w:jc w:val="both"/>')
    add_runs(p, text)


def bullets(items, num_id=3):
    for t in items:
        p = doc.add_paragraph(style="List Paragraph")
        set_ppr(p, f'<w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
                   '<w:spacing w:after="60"/><w:jc w:val="both"/>')
        add_runs(p, re.sub(r"^Step \d+:\s*", "", t))


def image_path(key: str) -> Path:
    if key.startswith("@"):
        k = key[1:]
        if k in RC.EXISTING:
            return ROOT / RC.EXISTING[k]
        fname, box = RC.SHOTS[k]
        im = Image.open(HERE / "screens" / fname).convert("RGB")
        if box:
            l, t, r, b = box
            im = im.crop((l, t, min(r, im.width), min(b, im.height)))
        out = CROPS / f"{k}.png"
        im.save(out)
        return out
    return HERE / "figures" / f"{key}.png"


# per-figure height caps (inches), tuned with work/gaps.py so tall screenshots pair up instead of leaving half pages
FIG_MAX_H: dict[str, float] = {"@pdf_chatgpt_p1": 8.3, "@pdf_chatgpt_p2": 8.3, "@pdf_video_p2": 8.3, "@roc_image": 2.3, "@video_explain": 7.4, "@audio_explain": 4.5, "@pdf_video": 4.1}


def figure(key: str, caption: str, width_in: float):
    global fig_n
    fig_n += 1
    path = image_path(key)
    w, h = Image.open(path).size
    max_h = FIG_MAX_H.get(key, 6.0)  # tall figures that will not fit below text leave half-empty pages
    if width_in * h / w > max_h:
        width_in = max_h * w / h
    p = doc.add_paragraph()
    set_ppr(p, '<w:keepNext/><w:spacing w:before="120" w:after="60" w:line="240" w:lineRule="auto"/><w:jc w:val="center"/>')
    p.add_run().add_picture(str(path), width=Inches(width_in))
    c = doc.add_paragraph(style="Figure Caption")
    r = c.add_run(f"Figure {label}.{fig_n}: "); r.bold = True
    add_runs(c, caption)


def table(caption: str, header, rows, widths):
    global tab_n
    tab_n += 1
    c = doc.add_paragraph(style="Table Caption")
    set_ppr(c, '<w:pStyle w:val="TableCaption"/><w:keepNext/>')
    r = c.add_run(f"Table {label}.{tab_n}: "); r.bold = True
    add_runs(c, caption)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl = t._tbl
    tbl.remove(tbl.find(qn("w:tblPr")))
    tblpr = copy.deepcopy(tpl_tblpr)
    tblpr.find(qn("w:tblW")).set(qn("w:w"), str(sum(widths)))
    tbl.insert(0, tblpr)
    grid = tbl.find(qn("w:tblGrid"))
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(wd))
    small = len(header) >= 5
    for ri, row in enumerate([header] + rows):
        tr = t.rows[ri]._tr
        if ri == 0:
            tr.insert(0, parse_xml(f'<w:trPr {nsdecls("w")}><w:tblHeader/><w:cantSplit/></w:trPr>'))
        else:
            tr.insert(0, parse_xml(f'<w:trPr {nsdecls("w")}><w:cantSplit/></w:trPr>'))
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            tc = cell._tc
            old = tc.find(qn("w:tcPr"))
            if old is not None:
                tc.remove(old)
            tcpr = copy.deepcopy(tpl_hdr_tcpr if ri == 0 else tpl_body_tcpr)
            tcpr.find(qn("w:tcW")).set(qn("w:w"), str(widths[ci]))
            tc.insert(0, tcpr)
            p = cell.paragraphs[0]
            align = "center" if (ri == 0 or (ci > 0 and len(str(val)) < 14)) else "left"
            set_ppr(p, f'<w:spacing w:after="0" w:line="240" w:lineRule="auto"/><w:jc w:val="{align}"/>')
            add_runs(p, str(val), size=9.5 if small else 10.5, bold_all=(ri == 0))
    sp = doc.add_paragraph()
    set_ppr(sp, '<w:spacing w:after="120"/>')


def _mt(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return f'<m:r><w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr><m:t xml:space="preserve">{text}</m:t></m:r>'


def omml(nodes) -> str:
    out = []
    for n in nodes:
        k = n[0]
        if k == "t":
            out.append(_mt(n[1]))
        elif k == "sub":
            out.append(f"<m:sSub><m:e>{_mt(n[1])}</m:e><m:sub>{_mt(n[2])}</m:sub></m:sSub>")
        elif k == "sup":
            out.append(f"<m:sSup><m:e>{_mt(n[1])}</m:e><m:sup>{_mt(n[2])}</m:sup></m:sSup>")
        elif k == "f":
            out.append(f"<m:f><m:num>{omml(n[1])}</m:num><m:den>{omml(n[2])}</m:den></m:f>")
        elif k == "nary":
            out.append(f'<m:nary><m:naryPr><m:chr m:val="{n[1]}"/></m:naryPr><m:sub>{_mt(n[2])}</m:sub>'
                       f"<m:sup>{_mt(n[3])}</m:sup><m:e>{omml(n[4])}</m:e></m:nary>")
    return "".join(out)


def equation(nodes, label: str):
    p = doc.add_paragraph()
    set_ppr(p, '<w:tabs><w:tab w:val="center" w:pos="4333"/><w:tab w:val="right" w:pos="8666"/></w:tabs>'
               '<w:spacing w:before="120" w:after="120" w:line="276" w:lineRule="auto"/>')
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:tab/></w:r>'))
    p._p.append(parse_xml(f'<m:oMath {M} {nsdecls("w")}>{omml(nodes)}</m:oMath>'))
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:tab/><w:t>({label})</w:t></w:r>'))


def appendix(letter: str, title: str):
    """Appendix title: unnumbered, but listed in the Contents (template style 'Front Matter Title (in TOC)')."""
    global label, fig_n, tab_n
    label, fig_n, tab_n = letter, 0, 0
    p = doc.add_paragraph(style="Front Matter Title (in TOC)")
    set_ppr(p, '<w:pStyle w:val="FrontMatterTitleinTOC"/><w:pageBreakBefore/>')
    p.add_run(f"Appendix {letter}: {title}")


def subhead(text: str):
    p = doc.add_paragraph()
    set_ppr(p, '<w:keepNext/><w:spacing w:before="240" w:after="120"/>')
    r = p.add_run(text)
    r.bold = True
    r.font.size = docx.shared.Pt(13)


def source_excerpt(rel: str, spec) -> str:
    """Verbatim excerpt of a project file: ("func", name) for one function or method (via ast), or
    ("between", start_text, end_text) for the lines from the first one containing start_text to the next one
    containing end_text, inclusive."""
    import ast
    import textwrap
    src = (ROOT / rel).read_text(encoding="utf-8")
    lines = src.splitlines()
    if spec[0] == "func":
        for node in ast.walk(ast.parse(src)):
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name == spec[1]:
                start = (node.decorator_list[0].lineno if node.decorator_list else node.lineno) - 1
                return textwrap.dedent(NL.join(lines[start:node.end_lineno]))
        raise KeyError(spec)
    i = next(n for n, line in enumerate(lines) if spec[1] in line)
    j = next(n for n, line in enumerate(lines[i + 1:], i + 1) if spec[2] in line)
    return textwrap.dedent(NL.join(lines[i:j + 1]))


NL = chr(10)


def _mono_box(lines, size: float):
    t = doc.add_table(rows=1, cols=1)
    tbl = t._tbl
    tbl.remove(tbl.find(qn("w:tblPr")))
    tbl.insert(0, parse_xml(
        f'<w:tblPr {nsdecls("w")}><w:tblW w:w="8666" w:type="dxa"/><w:tblBorders>'
        '<w:top w:val="single" w:sz="4" w:color="BFBFBF"/><w:left w:val="single" w:sz="4" w:color="BFBFBF"/>'
        '<w:bottom w:val="single" w:sz="4" w:color="BFBFBF"/><w:right w:val="single" w:sz="4" w:color="BFBFBF"/>'
        '</w:tblBorders><w:tblCellMar><w:top w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/>'
        '<w:bottom w:w="80" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tblCellMar></w:tblPr>'))
    tbl.find(qn("w:tblGrid")).find(qn("w:gridCol")).set(qn("w:w"), "8666")
    cell = t.rows[0].cells[0]
    cell._tc.insert(0, parse_xml(f'<w:tcPr {nsdecls("w")}><w:tcW w:w="8666" w:type="dxa"/>'
                                 '<w:shd w:val="clear" w:color="auto" w:fill="F4F4F2"/></w:tcPr>'))
    for n, line in enumerate(lines):
        p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        set_ppr(p, '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>')
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"
        r.font.size = docx.shared.Pt(size)
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Consolas")
    sp = doc.add_paragraph()
    set_ppr(sp, '<w:spacing w:after="120"/>')


def code(rel: str, spec, caption: str):
    cap = doc.add_paragraph()
    set_ppr(cap, '<w:keepNext/><w:spacing w:before="120" w:after="60"/>')
    r = cap.add_run(f"{caption} ")
    r.bold = True
    r = cap.add_run(f"({rel})")
    r.italic = True
    r.font.size = docx.shared.Pt(10)
    _mono_box(source_excerpt(rel, spec).split(NL), 8.5)


def cmd(lines):
    _mono_box(lines, 9)


# ------------------------------------------------------------------------------------------- render body
def render(blocks):
    for b in blocks:
        render_one(b)


def render_one(b):
    k = b[0]
    if k == "app":
        appendix(b[1], b[2])
    elif k == "ah":
        subhead(b[1])
    elif k == "code":
        code(b[1], b[2], b[3])
    elif k == "cmd":
        cmd(b[1])
    elif k == "h1":
        heading(1, b[1])
    elif k == "h2":
        heading(2, b[1])
    elif k == "h3":
        heading(3, b[1])
    elif k == "p":
        para(b[1])
    elif k == "bul":
        bullets(b[1])
    elif k == "num":
        bullets(b[1], new_num(1))
    elif k == "fig":
        figure(b[1], b[2], b[3])
    elif k == "tab":
        table(b[1], b[2], b[3], b[4])
    elif k == "eq":
        equation(b[1], b[2])
    else:
        raise ValueError(k)


render(RC.BODY)

# --- references (template style: ListParagraph, "[n]" numbering)
p = doc.add_paragraph(style="Front Matter Title (in TOC)")
set_ppr(p, '<w:pStyle w:val="FrontMatterTitleinTOC"/><w:pageBreakBefore/>')
p.add_run("References")
for key in ORDER:
    p = doc.add_paragraph(style="List Paragraph")
    p._p.insert(0, copy.deepcopy(ref_ppr))
    p.add_run(REFS[key])

render(getattr(RC, "APPENDICES", []))

# --- the empty paragraph that closes each TOC field (Contents, List of Figures, List of Tables) must take no space:
# when a list exactly fills its last page, that paragraph alone would otherwise spill onto a blank roman page
for sdt in body.iter(qn("w:sdt")):
    content = sdt.find(qn("w:sdtContent"))
    if content is None or not len(content):
        continue
    last = content[-1]
    texts = "".join(t.text or "" for t in last.iter(qn("w:t")))
    if last.tag == qn("w:p") and not texts.strip() and last.find(".//" + qn("w:fldChar")) is not None:
        old = last.find(qn("w:pPr"))
        if old is not None:
            last.remove(old)
        last.insert(0, parse_xml(f'<w:pPr {nsdecls("w")}><w:spacing w:before="0" w:after="0" w:line="20" '
                                 'w:lineRule="exact"/><w:rPr><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr></w:pPr>'))
        for r in last.findall(qn("w:r")):
            r.insert(0, parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr>'))

# --- ask Word to refresh the TOC / list of figures / list of tables on open
settings = doc.settings.element
if settings.find(qn("w:updateFields")) is None:
    settings.append(parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'))

doc.core_properties.title = RC.SHORT_TITLE
doc.core_properties.author = ", ".join(n for n, _ in RC.STUDENTS)
doc.save(str(OUT))
print("wrote", OUT, "| references:", len(ORDER), "| chapters:", chapter)
